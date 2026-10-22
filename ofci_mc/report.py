"""Excel workbook + PNG charts."""
import os
from typing import Dict

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.colors import LinearSegmentedColormap  # noqa: E402

# Reference palette (dataviz skill): first three categorical slots validate all-pairs.
SCENARIO_STYLE = {
    "Baseline": {"color": "#2a78d6", "ls": "-"},
    "Pre-mitigated": {"color": "#eb6834", "ls": "--"},
    "Post-mitigated": {"color": "#1baf7a", "ls": ":"},
}
INK, INK2, MUTED, GRID, AXIS, SURFACE = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7", "#fcfcfb"
SEQ_BLUE = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95", "#0d366b"]


def _style(ax):
    ax.set_facecolor(SURFACE)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(AXIS)
    ax.tick_params(colors=INK2, labelsize=8)
    ax.grid(True, color=GRID, linewidth=0.6)
    ax.set_axisbelow(True)


def _fig(w, h):
    fig = plt.figure(figsize=(w, h), dpi=150)
    fig.patch.set_facecolor(SURFACE)
    plt.rcParams["font.family"] = ["Helvetica", "Arial", "DejaVu Sans"]
    return fig


def chart_building_s_curves(bm: Dict[str, tuple], buildings: pd.DataFrame, as_of, path):
    keys = list(next(iter(bm.values()))[0])
    fig = _fig(11, 2.6 * ((len(keys) + 2) // 3))
    for i, b in enumerate(keys):
        ax = fig.add_subplot((len(keys) + 2) // 3, 3, i + 1)
        _style(ax)
        for scen, (k, Bm) in bm.items():
            v = np.sort(Bm[:, list(k).index(b)])
            d = as_of + pd.to_timedelta(v, unit="D")
            st = SCENARIO_STYLE[scen]
            ax.plot(d, np.linspace(0, 1, len(v)), color=st["color"], ls=st["ls"], lw=2, label=scen)
        row = buildings[(buildings["bldg"] == b) & (buildings["scenario"] == "Post-mitigated")].iloc[0]
        ax.axvline(row["compass_all_available"], color=MUTED, lw=1, ls="-")
        ax.text(row["compass_all_available"], 0.02, " Compass", color=INK2, fontsize=7, rotation=90, va="bottom")
        if pd.notna(row["latest_roj_complete"]):
            ax.axvline(row["latest_roj_complete"], color=INK, lw=1, ls="-.")
            ax.text(row["latest_roj_complete"], 0.98, " last ROJ complete", color=INK, fontsize=7, rotation=90, va="top")
        ax.axhline(0.8, color=MUTED, lw=0.8, ls=":")
        ax.set_title(f"{b} - all open OFCI available", fontsize=9, color=INK, loc="left")
        ax.set_ylim(0, 1)
        ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
        loc = mdates.AutoDateLocator(minticks=3, maxticks=5)
        ax.xaxis.set_major_locator(loc)
        ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(loc))
    handles, labels = ax.get_legend_handles_labels()
    fig.suptitle("Cumulative probability that every open OFCI unit in the building is available",
                 fontsize=10, color=INK, x=0.01, ha="left", y=0.995)
    fig.legend(handles, labels, loc="upper left", bbox_to_anchor=(0.01, 0.965), ncol=3, frameon=False, fontsize=8)
    fig.tight_layout(rect=(0, 0, 1, 0.93))
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def chart_units_missing(dist: Dict[str, np.ndarray], compass_value: float, path):
    fig = _fig(8, 3.6)
    ax = fig.add_subplot(111)
    _style(ax)
    lo = min(min(v.min() for v in dist.values()), compass_value)
    hi = max(v.max() for v in dist.values())
    bins = np.linspace(lo, hi, 60)
    for scen, v in dist.items():
        st = SCENARIO_STYLE[scen]
        h, e = np.histogram(v, bins=bins)
        ax.step(e[:-1], h / len(v), where="post", color=st["color"], ls=st["ls"], lw=2,
                label=f"{scen}  (P80 = {np.percentile(v, 80):,.0f})")
    ax.axvline(compass_value, color=MUTED, lw=1)
    ax.text(compass_value, ax.get_ylim()[1] * 0.95, f" Compass today: {compass_value:,.0f}", color=INK2, fontsize=8, va="top")
    ax.set_xlabel("Open units available after their ROJ install window closes", fontsize=8, color=INK2)
    ax.set_ylabel("Share of iterations", fontsize=8, color=INK2)
    ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    ax.legend(frameon=False, fontsize=8)
    ax.set_title("How many open units miss their ROJ window?", fontsize=10, color=INK, loc="left")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def chart_tornado(t: pd.DataFrame, path, top=15):
    t = t.head(top).iloc[::-1]
    fig = _fig(8, 0.32 * len(t) + 1.2)
    ax = fig.add_subplot(111)
    _style(ax)
    ax.grid(True, axis="x", color=GRID, linewidth=0.6)
    ax.grid(False, axis="y")
    is_redate = t["driver"].str.startswith("Compass re-date")
    colors = np.where(is_redate, "#2a78d6", "#eb6834")
    ax.barh(t["driver"].str.replace("Compass re-date: ", "Re-date: "), t["spearman_rho"], color=colors, height=0.6)
    for y, (r, c) in enumerate(zip(t["spearman_rho"], t["contribution_to_variance"])):
        ax.text(r + (0.01 if r >= 0 else -0.01), y, f"{r:.2f}  ({c:.0%})", va="center",
                ha="left" if r >= 0 else "right", fontsize=7, color=INK2)
    ax.axvline(0, color=AXIS, lw=1)
    ax.set_xlabel("Spearman rank correlation with total unit-days late vs ROJ (share of variance)", fontsize=8, color=INK2)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color="#2a78d6", label="Compass re-date drift (fitted)"),
                       Patch(color="#eb6834", label="Risk-register event")], frameon=False, fontsize=8, loc="lower right")
    ax.set_title("What drives schedule risk (Pre-mitigated)", fontsize=10, color=INK, loc="left")
    ax.set_xlim(min(0, t["spearman_rho"].min()) - 0.15, t["spearman_rho"].max() + 0.25)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def chart_heatmap(eq: pd.DataFrame, path):
    eq = eq.assign(share=eq["expected_units_missing_roj_window"] / eq["open_units"])
    m = eq.pivot(index="equipment", columns="bldg", values="share")
    m = m.loc[m.mean(axis=1).sort_values(ascending=False).index]
    fig = _fig(7, 0.38 * len(m) + 1.4)
    ax = fig.add_subplot(111)
    ax.set_facecolor(SURFACE)
    cmap = LinearSegmentedColormap.from_list("seq_blue", SEQ_BLUE)
    im = ax.imshow(m.to_numpy(float), cmap=cmap, vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(range(m.shape[1]), m.columns, fontsize=8, color=INK2)
    ax.set_yticks(range(m.shape[0]), m.index, fontsize=8, color=INK2)
    for s in ax.spines.values():
        s.set_visible(False)
    for i in range(m.shape[0]):
        for j in range(m.shape[1]):
            v = m.iat[i, j]
            if pd.notna(v):
                ax.text(j, i, f"{v:.0%}", ha="center", va="center", fontsize=7,
                        color="#ffffff" if v > 0.55 else INK)
    ax.set_xticks(np.arange(-.5, m.shape[1]), minor=True)
    ax.set_yticks(np.arange(-.5, m.shape[0]), minor=True)
    ax.grid(which="minor", color=SURFACE, linewidth=2)
    ax.tick_params(which="both", length=0)
    cb = fig.colorbar(im, ax=ax, fraction=0.04)
    cb.ax.tick_params(labelsize=7, colors=INK2)
    cb.outline.set_visible(False)
    ax.set_title("Share of open units expected to miss their ROJ window (Post-mitigated)",
                 fontsize=9, color=INK, loc="left")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def chart_backtest(bt: pd.DataFrame, path):
    fig = _fig(6, 3)
    ax = fig.add_subplot(111)
    _style(ax)
    ax.hist(bt["pit"], bins=10, range=(0, 1), color="#2a78d6", rwidth=0.92,
            weights=np.full(len(bt), 1 / len(bt)))
    ax.axhline(0.1, color=MUTED, ls="--", lw=1)
    ax.text(0.01, 0.105, "perfectly calibrated", fontsize=7, color=INK2)
    ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    ax.set_xlabel("Where the actual Aug-18 Compass date fell in the simulated range (PIT)", fontsize=8, color=INK2)
    ax.set_title(f"Holdout test: model fit before the window, scored on actual re-dates (n={len(bt):,} units, "
                 f"P10-P90 coverage {bt['inside_p10_p90'].mean():.0%})", fontsize=9, color=INK, loc="left")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def write_workbook(path: str, sheets: Dict[str, pd.DataFrame], charts: Dict[str, str]):
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
        wb = xw.book
        from openpyxl.drawing.image import Image
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill

        head_fill = PatternFill("solid", fgColor="184F95")
        for name, df in sheets.items():
            ws = wb[name[:31]]
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, start=1):
                c = ws.cell(row=1, column=j)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(wrap_text=True, vertical="top")
                width = max(10, min(48, int(max([len(str(col)) * 0.9] +
                                                [len(str(v)) for v in df[col].head(200)]) + 2)))
                ws.column_dimensions[c.column_letter].width = width
                letter = c.column_letter
                last = len(df) + 1
                if str(col).startswith("p_") or col in ("deterministic_percentile", "criticality_index",
                                                       "contribution_to_variance", "pit"):
                    for r in range(2, last + 1):
                        ws.cell(row=r, column=j).number_format = "0%"
                    if str(col).startswith("p_"):
                        ws.conditional_formatting.add(
                            f"{letter}2:{letter}{last}",
                            ColorScaleRule(start_type="num", start_value=0, start_color="D03B3B",
                                           mid_type="num", mid_value=0.5, mid_color="FAB219",
                                           end_type="num", end_value=1, end_color="0CA30C"))
                elif df[col].dtype.kind == "f":
                    for r in range(2, last + 1):
                        ws.cell(row=r, column=j).number_format = "#,##0.0"
            ws.row_dimensions[1].height = 42
            if len(df.columns):
                ws.auto_filter.ref = ws.dimensions
        if charts:
            ws = wb.create_sheet("Charts", 1)
            row = 1
            for title, img in charts.items():
                ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
                pic = Image(img)
                scale = 900 / pic.width
                pic.width, pic.height = pic.width * scale, pic.height * scale
                ws.add_image(pic, f"A{row + 1}")
                row += int(pic.height / 20) + 4
    return path


def ensure_dir(p):
    os.makedirs(p, exist_ok=True)
    return p
